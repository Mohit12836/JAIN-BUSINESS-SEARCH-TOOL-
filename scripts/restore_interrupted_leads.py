import os
import sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.config import get_master_excel_path

def restore_leads():
    excel_path = get_master_excel_path()
    if not os.path.exists(excel_path):
        print(f"Excel file not found at: {excel_path}")
        return

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    interrupted_keywords = [
        "browser target page",
        "target pag",
        "page.evaluate",
        "page.wait_for_timeout",
        "execution context",
        "portal redirected away"
    ]

    restored_count = 0
    restored_leads = []

    for r in range(2, ws.max_row + 1):
        status_val = str(ws.cell(row=r, column=26).value or "").strip()
        status_lower = status_val.lower()

        is_interrupted = any(kw in status_lower for kw in interrupted_keywords)
        if is_interrupted:
            lead_name = str(ws.cell(row=r, column=2).value or "").strip()
            restored_leads.append((r, lead_name, status_val))
            
            # Reset status back to Ready to Submit
            ws.cell(row=r, column=24, value="")  # JFJ_Business_ID
            ws.cell(row=r, column=25, value="")  # JFJ_Profile_URL
            ws.cell(row=r, column=26, value="Ready to Submit")  # Submission Status
            restored_count += 1

    if restored_count > 0:
        wb.save(excel_path)
        print(f"\n✅ SUCCESS: Successfully restored {restored_count} interrupted leads back to 'Ready to Submit'!")
        print("\nRestored leads preview:")
        for r, name, old_status in restored_leads[:15]:
            print(f" - Row {r}: {name} (Was: '{old_status}')")
        if len(restored_leads) > 15:
            print(f" ... and {len(restored_leads) - 15} more.")
    else:
        print("\nNo interrupted leads found to restore.")

if __name__ == "__main__":
    restore_leads()
