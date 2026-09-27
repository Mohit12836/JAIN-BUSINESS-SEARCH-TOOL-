import os
import sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from backend.config import get_master_excel_path

def check_duplicates():
    excel_path = get_master_excel_path()
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    submitted_names = set()
    for r in range(2, ws.max_row + 1):
        biz_id = str(ws.cell(r, 24).value or "").strip()
        st = str(ws.cell(r, 26).value or "").strip().lower()
        name = str(ws.cell(r, 2).value or "").strip().lower()
        if "submitted" in st or "live" in st or biz_id.startswith("JFJ-"):
            if name:
                submitted_names.add(name)

    print(f"Total Unique Submitted Business Names: {len(submitted_names)}")

    pending_leads = []
    conflict_leads = []
    truly_fresh_leads = []

    for r in range(2, ws.max_row + 1):
        st = str(ws.cell(r, 26).value or "").strip()
        st_lower = st.lower()
        name = str(ws.cell(r, 2).value or "").strip()
        name_lower = name.lower()

        if "ready to submit" in st_lower or st == "":
            pending_leads.append((r, name))
            if name_lower in submitted_names:
                conflict_leads.append((r, name))
            else:
                truly_fresh_leads.append((r, name))

    print(f"Total Pending in Sheet: {len(pending_leads)}")
    print(f"Conflict with already submitted names: {len(conflict_leads)}")
    print(f"Truly Fresh, Never-Submitted Businesses: {len(truly_fresh_leads)}")

    print("\nFirst 10 Truly Fresh Businesses:")
    for r, n in truly_fresh_leads[:10]:
        print(f" - Row {r}: {n}")

if __name__ == "__main__":
    check_duplicates()
