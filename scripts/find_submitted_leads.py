import os
import sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from backend.config import get_master_excel_path

def find_submitted():
    excel_path = get_master_excel_path()
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    submitted = []
    for r in range(2, ws.max_row + 1):
        biz_id = str(ws.cell(r, 24).value or "").strip()
        status = str(ws.cell(r, 26).value or "").strip()
        url = str(ws.cell(r, 25).value or "").strip()
        name = str(ws.cell(r, 2).value or "").strip()
        cat = str(ws.cell(r, 3).value or "").strip()
        phone = str(ws.cell(r, 5).value or "").strip()
        city = str(ws.cell(r, 9).value or "").strip()
        photo = str(ws.cell(r, 17).value or "").strip()
        if biz_id.startswith("JFJ-") or "Submitted" in status or "Live" in status:
            submitted.append({
                "row": r,
                "biz_id": biz_id,
                "name": name,
                "category": cat,
                "phone": phone,
                "city": city,
                "url": url,
                "status": status,
                "photo": photo
            })

    print(f"Total Submitted Leads in Sheet: {len(submitted)}")
    print("\nLast 15 Submitted Leads:")
    for s in submitted[-15:]:
        print(f"Row {s['row']} | ID: {s['biz_id']} | Name: {s['name']} | Status: {s['status']}")

if __name__ == "__main__":
    find_submitted()
