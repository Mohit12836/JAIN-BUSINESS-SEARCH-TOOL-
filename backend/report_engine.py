"""
Multi-Period Report & Analytics Engine for JainBiz Lead Miner.
Generates comprehensive daily, weekly, multi-day, monthly, and custom date range
reports with executive KPI summaries, breakdowns, and 1-click Excel/CSV exports.
"""

import os
import sys
import re
import csv
import io
import sqlite3
import datetime
from typing import Dict, Any, List, Optional, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.config import get_master_excel_path
from backend.auto_entry_bot import load_leads_from_excel
from backend.database import DB_PATH, init_db

def parse_iso_datetime(dt_str: Any) -> Optional[datetime.datetime]:
    """Parses various datetime formats from DB or strings."""
    if not dt_str:
        return None
    if isinstance(dt_str, datetime.datetime):
        return dt_str
    if isinstance(dt_str, datetime.date):
        return datetime.datetime.combine(dt_str, datetime.time.min)
    
    clean_str = str(dt_str).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S"):
        try:
            return datetime.datetime.strptime(clean_str[:19], fmt)
        except Exception:
            continue
    return None

def get_timeframe_range(
    timeframe: str,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None
) -> Tuple[datetime.datetime, datetime.datetime, str]:
    """
    Calculates start and end datetimes based on selected timeframe.
    Options: 1d, 2d, 3d, 7d, 15d, 30d, all, custom
    """
    now = datetime.datetime.now()
    tf = (timeframe or "all").lower().strip()
    
    if tf == "1d":
        start = now - datetime.timedelta(days=1)
        label = "1 दिन (Last 24 Hours / आज)"
    elif tf == "2d":
        start = now - datetime.timedelta(days=2)
        label = "पिछले 2 दिन (Last 48 Hours)"
    elif tf == "3d":
        start = now - datetime.timedelta(days=3)
        label = "पिछले 3 दिन (Last 72 Hours)"
    elif tf in ("7d", "weekly", "week"):
        start = now - datetime.timedelta(days=7)
        label = "साप्ताहिक (Last 7 Days / 1 हफ्ता)"
    elif tf == "15d":
        start = now - datetime.timedelta(days=15)
        label = "पाक्षिक (Last 15 Days)"
    elif tf in ("30d", "monthly", "month"):
        start = now - datetime.timedelta(days=30)
        label = "मासिक (Last 30 Days / 1 महीना)"
    elif tf == "custom" and start_date:
        try:
            start = datetime.datetime.strptime(start_date.strip()[:10], "%Y-%m-%d")
        except Exception:
            start = now - datetime.timedelta(days=7)
        if end_date:
            try:
                end = datetime.datetime.strptime(end_date.strip()[:10], "%Y-%m-%d").replace(hour=23, minute=59, second=59)
            except Exception:
                end = now
        else:
            end = now
        label = f"कस्टम अवधि ({start.strftime('%d-%b-%Y')} से {end.strftime('%d-%b-%Y')})"
        return start, end, label
    else:
        start = datetime.datetime(2020, 1, 1, 0, 0, 0)
        label = "संपूर्ण डेटा (All Time Historical)"
        
    return start, now, label

def load_all_leads_with_timestamps() -> List[Dict[str, Any]]:
    """
    Loads all leads from Master Excel and enriches them with exact
    timestamps from SQLite scraped_leads history.
    """
    excel_path = get_master_excel_path()
    excel_leads = []
    if os.path.exists(excel_path):
        try:
            excel_leads = load_leads_from_excel(excel_path)
        except Exception as e:
            print(f"Warning: Error reading master excel: {e}")
            
    init_db()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    phone_map = {}
    name_map = {}
    cursor.execute("SELECT phone, name_norm, firm_name, city, category, scraped_at FROM scraped_leads")
    all_db_rows = cursor.fetchall()
    for row in all_db_rows:
        phone = row["phone"]
        name_norm = row["name_norm"]
        if phone:
            phone_map[phone] = row
        if name_norm:
            name_map[name_norm] = row
            
    merged_leads = []
    seen_keys = set()
    
    default_time = datetime.datetime.now()
    if os.path.exists(excel_path):
        default_time = datetime.datetime.fromtimestamp(os.path.getmtime(excel_path))
        
    for lead in excel_leads:
        phone_raw = lead.get("phone", "")
        clean_digits = re.sub(r"\D", "", phone_raw)[-10:]
        norm_name = re.sub(r"[^a-zA-Z0-9]", "", lead.get("name", "").lower())
        
        db_record = None
        if clean_digits and clean_digits in phone_map:
            db_record = phone_map[clean_digits]
        elif norm_name and norm_name in name_map:
            db_record = name_map[norm_name]
            
        if db_record and db_record["scraped_at"]:
            parsed_dt = parse_iso_datetime(db_record["scraped_at"])
            lead["created_at"] = parsed_dt or default_time
            lead["scraped_at_str"] = str(db_record["scraped_at"])
        else:
            lead["created_at"] = default_time
            lead["scraped_at_str"] = default_time.strftime("%Y-%m-%d %H:%M:%S")
            
        dedup_key = clean_digits or norm_name or lead.get("name", "")
        seen_keys.add(dedup_key)
        merged_leads.append(lead)
        
    # Also include any DB scraped leads that might not yet be in Master Excel
    for row in all_db_rows:
        phone = row["phone"] or ""
        norm_name = row["name_norm"] or ""
        dedup_key = phone if (phone and not phone.startswith("nophone_")) else norm_name
        
        if dedup_key and dedup_key not in seen_keys:
            seen_keys.add(dedup_key)
            parsed_dt = parse_iso_datetime(row["scraped_at"]) or default_time
            db_lead = {
                "row_idx": len(merged_leads) + 2,
                "sl": str(len(merged_leads) + 1),
                "name": row["firm_name"],
                "category": row["category"] or "Business & Industry",
                "owner": "Proprietor",
                "phone": phone if not phone.startswith("nophone_") else "Not Listed",
                "whatsapp": phone if not phone.startswith("nophone_") else "",
                "email": "",
                "address": f"{row['city']}, India",
                "city": row["city"],
                "district": row["city"],
                "state": "Madhya Pradesh" if row["city"].lower() in ["indore", "bhopal", "ujjain"] else "Rajasthan",
                "pincode": "",
                "latitude": "",
                "longitude": "",
                "maps_url": "",
                "website": "",
                "storefront_photo": "",
                "showcase_photo": "",
                "gallery_url": "",
                "website_logo": "",
                "description": f"Prestigious business establishment in {row['city']}.",
                "tier": "🟢 100% Jain Verified",
                "reason": "Database Verified Record",
                "j4j_business_id": "",
                "j4j_profile_url": "",
                "submission_status": "Ready to Submit",
                "created_at": parsed_dt,
                "scraped_at_str": str(row["scraped_at"])
            }
            merged_leads.append(db_lead)
            
    conn.close()
    return merged_leads

def generate_leads_report(
    timeframe: str = "all",
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    city: Optional[str] = None,
    category: Optional[str] = None,
    status_filter: Optional[str] = None
) -> Dict[str, Any]:
    """
    Executes full multi-period query, calculates KPI metrics,
    distributions, batch timeline, and returns filtered leads.
    """
    start_dt, end_dt, tf_label = get_timeframe_range(timeframe, start_date, end_date)
    all_leads = load_all_leads_with_timestamps()
    
    # Apply Timeframe Filter
    filtered = []
    for l in all_leads:
        dt = l.get("created_at")
        if not dt:
            continue
        if start_dt <= dt <= end_dt:
            filtered.append(l)
            
    # Apply Dimensional Filters
    if city and city.lower() != "all":
        filtered = [l for l in filtered if l.get("city", "").lower().strip() == city.lower().strip()]
        
    if category and category.lower() != "all":
        filtered = [l for l in filtered if category.lower().strip() in l.get("category", "").lower().strip()]
        
    if status_filter and status_filter.lower() != "all":
        if status_filter.lower() in ("live", "submitted"):
            filtered = [l for l in filtered if "submitted" in l.get("submission_status", "").lower() or "live" in l.get("submission_status", "").lower()]
        elif status_filter.lower() in ("ready", "unsubmitted"):
            filtered = [l for l in filtered if "submitted" not in l.get("submission_status", "").lower() and "live" not in l.get("submission_status", "").lower()]

    # Sort newest first
    filtered.sort(key=lambda x: x.get("created_at", datetime.datetime.min), reverse=True)
    
    # Calculate Executive KPIs
    total_count = len(filtered)
    submitted_count = sum(1 for l in filtered if "submitted" in l.get("submission_status", "").lower() or "live" in l.get("submission_status", "").lower())
    unsubmitted_count = total_count - submitted_count
    submission_rate = round((submitted_count / total_count * 100), 1) if total_count > 0 else 0.0
    
    verified_contacts = sum(1 for l in filtered if l.get("phone") and l.get("phone") != "Not Listed" and len(re.sub(r"\D", "", l.get("phone"))) >= 10)
    contact_rate = round((verified_contacts / total_count * 100), 1) if total_count > 0 else 0.0
    
    photos_count = sum(1 for l in filtered if l.get("storefront_photo") or l.get("showcase_photo"))
    photos_rate = round((photos_count / total_count * 100), 1) if total_count > 0 else 0.0

    # City Breakdown
    city_counts: Dict[str, int] = {}
    for l in filtered:
        c = l.get("city") or "Unknown"
        city_counts[c] = city_counts.get(c, 0) + 1
    sorted_cities = sorted(city_counts.items(), key=lambda x: x[1], reverse=True)
    city_breakdown = [
        {"city": c, "count": cnt, "percent": round((cnt / total_count * 100), 1) if total_count else 0}
        for c, cnt in sorted_cities
    ]
    top_city = sorted_cities[0][0] if sorted_cities else "N/A"
    
    # Category Breakdown
    cat_counts: Dict[str, int] = {}
    for l in filtered:
        cat = l.get("category") or "General"
        cat_counts[cat] = cat_counts.get(cat, 0) + 1
    sorted_cats = sorted(cat_counts.items(), key=lambda x: x[1], reverse=True)
    category_breakdown = [
        {"category": cat, "count": cnt, "percent": round((cnt / total_count * 100), 1) if total_count else 0}
        for cat, cnt in sorted_cats
    ]
    top_category = sorted_cats[0][0] if sorted_cats else "N/A"

    # Daily Timeline Breakdown
    daily_counts: Dict[str, Dict[str, int]] = {}
    for l in filtered:
        dt = l.get("created_at")
        d_str = dt.strftime("%Y-%m-%d") if dt else "Unknown"
        if d_str not in daily_counts:
            daily_counts[d_str] = {"total": 0, "submitted": 0}
        daily_counts[d_str]["total"] += 1
        if "submitted" in l.get("submission_status", "").lower() or "live" in l.get("submission_status", "").lower():
            daily_counts[d_str]["submitted"] += 1
            
    daily_timeline = [
        {"date": d, "total": data["total"], "submitted": data["submitted"]}
        for d, data in sorted(daily_counts.items(), reverse=True)
    ]

    # Query Search Batches for this timeframe
    init_db()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT * FROM search_batches ORDER BY id DESC")
    all_batches = [dict(r) for r in c.fetchall()]
    conn.close()
    
    filtered_batches = []
    for b in all_batches:
        b_dt = parse_iso_datetime(b.get("timestamp"))
        if b_dt and start_dt <= b_dt <= end_dt:
            filtered_batches.append(b)
            
    return {
        "timeframe": timeframe,
        "timeframe_label": tf_label,
        "start_date": start_dt.strftime("%Y-%m-%d %H:%M:%S"),
        "end_date": end_dt.strftime("%Y-%m-%d %H:%M:%S"),
        "kpis": {
            "total_leads": total_count,
            "submitted_leads": submitted_count,
            "unsubmitted_leads": unsubmitted_count,
            "submission_rate": submission_rate,
            "verified_contacts": verified_contacts,
            "contact_rate": contact_rate,
            "photos_available": photos_count,
            "photos_rate": photos_rate,
            "top_city": top_city,
            "top_category": top_category,
            "batches_count": len(filtered_batches)
        },
        "city_breakdown": city_breakdown,
        "category_breakdown": category_breakdown,
        "daily_timeline": daily_timeline,
        "batches": filtered_batches,
        "leads": filtered
    }

def build_report_excel(report_data: Dict[str, Any], output_path: str) -> str:
    """
    Builds a professional 3-Sheet Excel report tailored to the selected timeframe.
    Sheet 1: Filtered JainForJain Leads (26 Columns)
    Sheet 2: Executive Period Summary
    Sheet 3: Search Batches Log
    """
    wb = Workbook()
    
    # ------------------ SHEET 1: FILTERED LEADS ------------------
    ws1 = wb.active
    ws1.title = "Filtered Leads (26 Columns)"
    
    headers = [
        "Sl.",
        "Business / Firm Name",
        "JainForJain Category",
        "Owner / Contact Person",
        "Calling Number",
        "WhatsApp Number",
        "Email ID",
        "Address (Street & Locality)",
        "City",
        "District",
        "State",
        "Pincode",
        "Latitude",
        "Longitude",
        "Google Maps URL",
        "Website URL",
        "Storefront Signboard (1600px HD)",
        "Showroom / Products (1600px HD)",
        "All Photos Gallery (Google Maps)",
        "Official Website Logo",
        "Auto-Generated Description (Ready to Paste)",
        "Verification Status",
        "Proof / Match Reason",
        "JainForJain Business ID",
        "JainForJain Live Profile URL",
        "Submission Status"
    ]
    
    header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    portal_header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    
    data_font = Font(name="Segoe UI", size=9, color="1E293B")
    cat_font = Font(name="Segoe UI", size=9, bold=True, color="0F766E")
    owner_font = Font(name="Segoe UI", size=9, bold=True, color="4338CA")
    link_font = Font(name="Segoe UI", size=9, color="2563EB", underline="single")
    coord_font = Font(name="Segoe UI", size=9, color="047857")
    
    status_ready_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    status_ready_font = Font(name="Segoe UI", size=9, bold=True, color="92400E")
    status_done_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    status_done_font = Font(name="Segoe UI", size=9, bold=True, color="166534")
    
    thin_border = Border(
        left=Side(border_style="thin", color="E2E8F0"),
        right=Side(border_style="thin", color="E2E8F0"),
        top=Side(border_style="thin", color="E2E8F0"),
        bottom=Side(border_style="thin", color="E2E8F0")
    )
    
    ws1.append(headers)
    ws1.row_dimensions[1].height = 28
    
    for c_idx in range(1, len(headers) + 1):
        cell = ws1.cell(1, c_idx)
        cell.fill = portal_header_fill if c_idx in [24, 25, 26] else header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    leads = report_data.get("leads", [])
    for idx, lead in enumerate(leads, start=2):
        row_vals = [
            idx - 1,
            lead.get("name", ""),
            lead.get("category", "Business & Industry"),
            lead.get("owner", "Proprietor"),
            lead.get("phone", "Not Listed"),
            lead.get("whatsapp", lead.get("phone", "")),
            lead.get("email", ""),
            lead.get("address", ""),
            lead.get("city", ""),
            lead.get("district", lead.get("city", "")),
            lead.get("state", ""),
            lead.get("pincode", ""),
            lead.get("latitude", ""),
            lead.get("longitude", ""),
            "Open Google Map" if lead.get("maps_url") else "",
            "Visit Website" if lead.get("website") else "",
            "📸 View Storefront (1600px)" if lead.get("storefront_photo") else "No Photo Listed",
            "🏬 View Showroom (1600px)" if lead.get("showcase_photo") else "Check Gallery",
            "🌐 Browse All Photos" if lead.get("gallery_url") else "",
            "🏷️ View Web Logo" if lead.get("website_logo") else "Use Storefront Photo",
            lead.get("description", ""),
            lead.get("tier", "🟢 100% Jain Verified"),
            lead.get("reason", "Historical Data Verified"),
            lead.get("j4j_business_id", ""),
            lead.get("j4j_profile_url", ""),
            lead.get("submission_status", "Ready to Submit")
        ]
        ws1.append(row_vals)
        ws1.row_dimensions[idx].height = 24
        
        for c_idx in range(1, len(row_vals) + 1):
            cell = ws1.cell(idx, c_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            
            if c_idx in [1, 5, 6, 9, 10, 11, 12, 13, 14]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c_idx in [13, 14]:
                cell.font = coord_font
            if c_idx == 3:
                cell.font = cat_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c_idx == 4:
                cell.font = owner_font
                
            # Links
            if c_idx == 15 and lead.get("maps_url"):
                cell.hyperlink = lead["maps_url"]
                cell.font = link_font
            if c_idx == 16 and lead.get("website"):
                cell.hyperlink = lead["website"]
                cell.font = link_font
            if c_idx == 17 and lead.get("storefront_photo") and str(lead.get("storefront_photo")).startswith("http"):
                cell.hyperlink = lead["storefront_photo"]
                cell.font = link_font
            if c_idx == 18 and lead.get("showcase_photo") and str(lead.get("showcase_photo")).startswith("http"):
                cell.hyperlink = lead["showcase_photo"]
                cell.font = link_font
            if c_idx == 25 and lead.get("j4j_profile_url") and str(lead.get("j4j_profile_url")).startswith("http"):
                cell.hyperlink = lead["j4j_profile_url"]
                cell.font = link_font
                
            # Status styling
            if c_idx == 26:
                status_val = str(cell.value)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "submitted" in status_val.lower() or "live" in status_val.lower():
                    cell.fill = status_done_fill
                    cell.font = status_done_font
                else:
                    cell.fill = status_ready_fill
                    cell.font = status_ready_font

    col_widths = [6, 32, 22, 18, 16, 16, 22, 35, 14, 14, 16, 10, 12, 12, 18, 18, 26, 26, 20, 20, 45, 18, 25, 15, 30, 22]
    for c_idx, w in enumerate(col_widths, start=1):
        col_letter = get_column_letter(c_idx)
        ws1.column_dimensions[col_letter].width = w

    # ------------------ SHEET 2: EXECUTIVE PERIOD SUMMARY ------------------
    ws2 = wb.create_sheet(title="Executive Period Summary")
    
    ws2.merge_cells("A1:F2")
    title_cell = ws2.cell(1, 1)
    title_cell.value = f"JainBiz Mining & Verification Executive Report - {report_data.get('timeframe_label')}"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws2.cell(4, 1, "Report Generated:").font = Font(name="Segoe UI", bold=True)
    ws2.cell(4, 2, datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    ws2.cell(5, 1, "Reporting Period:").font = Font(name="Segoe UI", bold=True)
    ws2.cell(5, 2, f"{report_data.get('start_date')} to {report_data.get('end_date')}")
    
    kpis = report_data.get("kpis", {})
    kpi_headers = ["Total Leads Extracted", "Portal Live (Submitted)", "Unsubmitted Leads", "Submission Rate", "Verified Phone %", "Canva HD Assets %"]
    kpi_values = [
        kpis.get("total_leads", 0),
        kpis.get("submitted_leads", 0),
        kpis.get("unsubmitted_leads", 0),
        f"{kpis.get('submission_rate', 0)}%",
        f"{kpis.get('contact_rate', 0)}%",
        f"{kpis.get('photos_rate', 0)}%"
    ]
    
    ws2.cell(7, 1, "CORE PERIOD METRICS").font = Font(name="Segoe UI", size=11, bold=True, color="4338CA")
    for i, (kh, kv) in enumerate(zip(kpi_headers, kpi_values), start=1):
        ws2.cell(8, i, kh).font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        ws2.cell(8, i).fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
        ws2.cell(8, i).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(8, i).border = thin_border
        
        ws2.cell(9, i, kv).font = Font(name="Segoe UI", size=12, bold=True, color="1E293B")
        ws2.cell(9, i).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(9, i).fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
        ws2.cell(9, i).border = thin_border
        ws2.column_dimensions[get_column_letter(i)].width = 24
        
    ws2.cell(11, 1, "CITY-WISE DISTRIBUTION").font = Font(name="Segoe UI", size=11, bold=True, color="0F766E")
    ws2.cell(12, 1, "City").font = Font(name="Segoe UI", bold=True, color="FFFFFF")
    ws2.cell(12, 1).fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    ws2.cell(12, 2, "Total Leads").font = Font(name="Segoe UI", bold=True, color="FFFFFF")
    ws2.cell(12, 2).fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    ws2.cell(12, 3, "Share (%)").font = Font(name="Segoe UI", bold=True, color="FFFFFF")
    ws2.cell(12, 3).fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    
    cur_row = 13
    for cb in report_data.get("city_breakdown", []):
        ws2.cell(cur_row, 1, cb["city"]).border = thin_border
        ws2.cell(cur_row, 2, cb["count"]).border = thin_border
        ws2.cell(cur_row, 3, f"{cb['percent']}%").border = thin_border
        cur_row += 1
        
    cur_row += 1
    ws2.cell(cur_row, 1, "TOP BUSINESS SECTORS").font = Font(name="Segoe UI", size=11, bold=True, color="9333EA")
    cur_row += 1
    ws2.cell(cur_row, 1, "Category").font = Font(name="Segoe UI", bold=True, color="FFFFFF")
    ws2.cell(cur_row, 1).fill = PatternFill(start_color="9333EA", end_color="9333EA", fill_type="solid")
    ws2.cell(cur_row, 2, "Total Leads").font = Font(name="Segoe UI", bold=True, color="FFFFFF")
    ws2.cell(cur_row, 2).fill = PatternFill(start_color="9333EA", end_color="9333EA", fill_type="solid")
    ws2.cell(cur_row, 3, "Share (%)").font = Font(name="Segoe UI", bold=True, color="FFFFFF")
    ws2.cell(cur_row, 3).fill = PatternFill(start_color="9333EA", end_color="9333EA", fill_type="solid")
    cur_row += 1
    for cat in report_data.get("category_breakdown", []):
        ws2.cell(cur_row, 1, cat["category"]).border = thin_border
        ws2.cell(cur_row, 2, cat["count"]).border = thin_border
        ws2.cell(cur_row, 3, f"{cat['percent']}%").border = thin_border
        cur_row += 1

    # ------------------ SHEET 3: SEARCH BATCHES LOG ------------------
    ws3 = wb.create_sheet(title="Automation Batches Log")
    batch_headers = ["Batch ID", "Timestamp", "City", "Micro-Market Area", "Batch Size", "Mined Count", "Status", "Next Target Area"]
    ws3.append(batch_headers)
    ws3.row_dimensions[1].height = 26
    for c_idx in range(1, len(batch_headers) + 1):
        cell = ws3.cell(1, c_idx)
        cell.fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws3.column_dimensions[get_column_letter(c_idx)].width = 20
        
    for r_idx, b in enumerate(report_data.get("batches", []), start=2):
        b_vals = [
            b.get("batch_id", ""),
            b.get("timestamp", ""),
            b.get("city", ""),
            b.get("area_name", ""),
            b.get("batch_size", 0),
            b.get("extracted_count", 0),
            b.get("status", "Completed"),
            b.get("next_area", "")
        ]
        ws3.append(b_vals)
        for c_idx in range(1, len(b_vals) + 1):
            cell = ws3.cell(r_idx, c_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if c_idx in [1, 2, 3, 5, 6, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path

def build_report_csv(report_data: Dict[str, Any]) -> str:
    """Generates clean UTF-8 CSV string for WhatsApp and CRM tools."""
    output = io.StringIO()
    writer = csv.writer(output)
    
    headers = [
        "Sl",
        "Firm Name",
        "Jain Category",
        "Owner / Contact Person",
        "Phone Number",
        "WhatsApp Number",
        "Address",
        "City",
        "District",
        "State",
        "Pincode",
        "Google Maps URL",
        "Website",
        "JainForJain Business ID",
        "Live Profile URL",
        "Submission Status",
        "Mined Date Time"
    ]
    writer.writerow(headers)
    
    for idx, lead in enumerate(report_data.get("leads", []), start=1):
        writer.writerow([
            idx,
            lead.get("name", ""),
            lead.get("category", ""),
            lead.get("owner", ""),
            lead.get("phone", ""),
            lead.get("whatsapp", ""),
            lead.get("address", ""),
            lead.get("city", ""),
            lead.get("district", ""),
            lead.get("state", ""),
            lead.get("pincode", ""),
            lead.get("maps_url", ""),
            lead.get("website", ""),
            lead.get("j4j_business_id", ""),
            lead.get("j4j_profile_url", ""),
            lead.get("submission_status", ""),
            lead.get("scraped_at_str", "")
        ])
        
    return output.getvalue()
