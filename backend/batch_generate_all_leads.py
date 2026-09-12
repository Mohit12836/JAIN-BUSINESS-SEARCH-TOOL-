"""
Batch Canva Asset Generator & Sync for all leads in Master Excel.
Renders 1080x1080 Logo and 1200x500 Banner for every firm in data/Jain_Leads_Verified_Photos_HD.xlsx.
Updates Excel links and syncs to Google Sheets.
"""

import os
import sys
import asyncio
import openpyxl
from openpyxl.styles import Font, Alignment

# Ensure backend directory is in path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.canva_storefront_generator import generate_batch_assets, get_firm_asset_slug, OUTPUT_DIR
from backend.auto_entry_bot import read_leads_from_excel
from backend.google_sheets_sync import sync_excel_to_google_sheet
from backend.config import get_master_excel_path

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

async def main():
    excel_path = get_master_excel_path()
    repo_excel = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "Jain_Leads_Verified_Photos_HD.xlsx")
    desktop_excel = os.path.join(os.path.expanduser("~"), "Desktop", "Jain_Leads_Verified_Photos_HD.xlsx")

    print(f"📖 Reading leads from: {excel_path}")
    leads = read_leads_from_excel(excel_path)
    print(f"Found {len(leads)} leads in Master Excel.")

    # 1. Batch generate assets
    print("\n🎨 Generating Canva Pro assets (1080x1080 Logo & 1200x500 Banner) for all leads...")
    leads = await generate_batch_assets(leads, max_concurrent=4, force=True)

    # 2. Update Excel files
    targets = [repo_excel]
    if os.path.exists(os.path.dirname(desktop_excel)):
        targets.append(desktop_excel)

    link_font = Font(name="Segoe UI", size=9, color="2563EB", underline="single")

    for target_path in set(targets):
        if not os.path.exists(target_path):
            continue
        print(f"\n💾 Updating links in: {target_path}")
        wb = openpyxl.load_workbook(target_path)
        ws = wb.active

        for lead in leads:
            row_idx = lead["row_idx"]
            slug = get_firm_asset_slug(lead["name"])
            banner_url = f"http://127.0.0.1:8000/api/canva-asset/{slug}_banner_1200x500.png"
            logo_url = f"http://127.0.0.1:8000/api/canva-asset/{slug}_logo_1080x1080.png"

            # Col 17: Storefront Signboard Banner
            c17 = ws.cell(row=row_idx, column=17)
            c17.value = "🎨 Canva Pro Storefront Banner"
            c17.hyperlink = banner_url
            c17.font = link_font
            c17.alignment = Alignment(horizontal="center", vertical="center")

            # Col 18: Showcase / Logo
            c18 = ws.cell(row=row_idx, column=18)
            c18.value = "💎 Canva Pro Profile Logo"
            c18.hyperlink = logo_url
            c18.font = link_font
            c18.alignment = Alignment(horizontal="center", vertical="center")

            # Col 19: All Photos Gallery (Google Maps)
            maps_url = lead.get("maps_url", "")
            c19 = ws.cell(row=row_idx, column=19)
            c19.value = "🌐 Browse All Photos"
            if maps_url:
                c19.hyperlink = maps_url
            c19.font = link_font
            c19.alignment = Alignment(horizontal="center", vertical="center")

            # Col 20: Official Website Logo
            c20 = ws.cell(row=row_idx, column=20)
            c20.value = "🏷️ Canva Pro Official Logo"
            c20.hyperlink = logo_url
            c20.font = link_font
            c20.alignment = Alignment(horizontal="center", vertical="center")

        wb.save(target_path)
        print(f"✓ Saved {len(leads)} updated leads to {target_path}")

    # 3. Verify files on disk
    print("\n🔍 Verifying all Canva asset files in data/canva_storefronts/:")
    all_ok = True
    for lead in leads:
        slug = get_firm_asset_slug(lead["name"])
        b_file = os.path.join(OUTPUT_DIR, f"{slug}_banner_1200x500.png")
        l_file = os.path.join(OUTPUT_DIR, f"{slug}_logo_1080x1080.png")
        b_exists = os.path.exists(b_file) and os.path.getsize(b_file) > 10000
        l_exists = os.path.exists(l_file) and os.path.getsize(l_file) > 10000
        status = "✅ OK" if (b_exists and l_exists) else "❌ MISSING"
        if not (b_exists and l_exists):
            all_ok = False
        print(f"[{status}] {lead['name'][:30]:<30} -> Banner: {os.path.basename(b_file)} ({os.path.getsize(b_file) if b_exists else 0} B) | Logo: {os.path.basename(l_file)} ({os.path.getsize(l_file) if l_exists else 0} B)")

    print(f"\nAsset Generation Status: {'100% COMPLETE' if all_ok else 'SOME MISSING'}")

    # 4. Sync to Google Sheets
    print("\n📊 Syncing to Google Sheets...")
    try:
        ok = await sync_excel_to_google_sheet(excel_path)
        print(f"Google Sheets Sync: {'✅ SUCCESS' if ok else '⚠️ FAILED'}")
    except Exception as ge:
        print(f"Google Sheets sync error: {ge}")

if __name__ == "__main__":
    asyncio.run(main())
